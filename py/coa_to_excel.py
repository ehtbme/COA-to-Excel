#!/usr/bin/env python3
"""Build a chemical inventory Excel sheet from a zip of COA (Certificate of Analysis) PDFs.

For every PDF in the zip, the first page is read and these columns are filled:
Chemical, Synonyms, Lot No., Product No., Expiry Date, CAS No., Manufacturer,
Form, Concentration, Chemical Formula, Drug Mol. Wt.
A value that the COA does not provide is written as "NA". All other columns are
left blank for manual entry.

Each COA is also saved as <LotNo>_<ExpiryDate>.pdf (e.g. 0612766_09JUN2029.pdf).

Usage:
    python coa_to_excel.py COAs.zip [more.zip ...] [-o OUTPUT_DIR]
"""

import argparse
import io
import re
import sys
import zipfile
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = [
    "Bottle no.", "Date Received", "Chemical", "Synonyms", "Lot No.", "Product No.",
    "Expiry Date", "CAS No.", "Manufacturer", "Container", "Cap Color", "Amount", "Form",
    "Class", "Concentration", "Location", "Chemical Formula", "Drug Mol. Wt.",
    "Name of Officer", "Remarks",
]
NA = "NA"

MONTHS = {m: i for i, m in enumerate(
    ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"], 1)}
MONTH_ABBR = {v: k for k, v in MONTHS.items()}

KNOWN_MANUFACTURERS = [
    ("cayman", "Cayman Chemical"),
    ("lgc", "LGC Standards"),
    ("cerilliant", "Cerilliant"),
    ("sigma-aldrich", "Sigma-Aldrich"),
    ("lipomed", "Lipomed"),
    ("chiron", "Chiron"),
    ("toronto research", "Toronto Research Chemicals"),
    ("restek", "Restek"),
]

# Soft hyphen / unicode dashes -> "-"
DASHES = str.maketrans({"­": "-", "‐": "-", "‑": "-", "‒": "-", "–": "-"})


# --------------------------------------------------------------------------- PDF text

def page_lines(page):
    """Return the page's text lines, with sub/superscripts kept on their parent line.

    pdfplumber's default extraction puts subscripts (C11H15NO, -d5) on a separate line.
    Here smaller characters that touch a larger-font line are merged into it, and
    superscripts are prefixed with '^' so isotope labels (^2H) can be recognised.
    """
    chars = [c for c in page.dedupe_chars().chars if c["text"].strip()]
    chars.sort(key=lambda c: (c["top"], c["x0"]))

    # 1. group characters with (almost) the same top into lines
    lines = []
    for c in chars:
        if lines and abs(c["top"] - lines[-1][0]["top"]) <= 2:
            lines[-1].append(c)
        else:
            lines.append([c])

    def box(line):
        return (min(c["top"] for c in line), max(c["bottom"] for c in line),
                max(c["size"] for c in line))

    # 2. move smaller digits onto an adjacent larger-font line they overlap with
    for i, src in enumerate(lines):
        for c in sorted(src, key=lambda c: c["x0"]):
            if not c["text"].isdigit():  # sub/superscripts on COAs are numbers
                continue
            for j, host in enumerate(lines):
                if i == j or not host:
                    continue
                top, bottom, size = box(host)
                if size < c["size"] * 1.15 or not (c["top"] < bottom and c["bottom"] > top):
                    continue
                gap = min(max(0, max(h["x0"], c["x0"]) - min(h["x1"], c["x1"])) for h in host)
                if gap <= 0.5 * size:
                    host.append(c)
                    src.remove(c)
                    break

    # 3. build strings; a wide horizontal gap becomes 3 spaces (column separator)
    out = []
    for line in lines:
        if not line:
            continue
        line.sort(key=lambda c: c["x0"])
        size = max(c["size"] for c in line)
        main = [c for c in line if c["size"] == size]
        mid = sum((c["top"] + c["bottom"]) / 2 for c in main) / len(main)
        for c in line:  # small digit sitting above the line's middle = superscript
            c["_sup"] = (c["text"].isdigit() and c["size"] < 0.87 * size
                         and (c["top"] + c["bottom"]) / 2 < mid - 0.1 * size)
        text, prev = "", None
        for c in line:
            if prev is not None:
                gap = c["x0"] - prev["x1"]
                big = max(c["size"], prev["size"])
                if gap > 2 * big:
                    text += "   "
                elif gap > 0.15 * big:
                    text += " "
            text += ("^" if c["_sup"] else "") + c["text"]
            prev = c
        out.append({"text": text.translate(DASHES), "x1": line[-1]["x1"]})
    return out


# --------------------------------------------------------------------------- parsing helpers

def search(pattern, text, flags=re.I | re.M):
    m = re.search(pattern, text, flags)
    return m.group(1).strip() if m else None


def plain(s):
    """Drop superscript markers for display (e.g. Δ^9-THC -> Δ9-THC)."""
    return s.replace("^", "") if s else s


def first_segment(s):
    """Text before the first wide gap (drops right-hand columns / structure labels)."""
    return re.split(r"\s{3,}", s.strip())[0].strip()


def format_date(text):
    """Convert a date string to DDMMMYYYY; day defaults to 01 when not given."""
    if not text:
        return None
    t = text.strip()
    patterns = [  # (regex, order of day/month/year groups)
        (r"(\d{1,2})\s*[-/ .]?\s*([A-Za-z]{3,9})\.?\s*[-/ ,.]?\s*(\d{4})", "dmy"),  # 09JUN2029, 30 Apr 2025
        (r"([A-Za-z]{3,9})\.?\s+(\d{1,2}),?\s+(\d{4})", "mdy"),                    # June 9, 2029
        (r"(\d{4})-(\d{1,2})-(\d{1,2})", "ymd"),                                    # 2029-06-09
        (r"(\d{1,2})[/.](\d{1,2})[/.](\d{4})", "dmy"),                              # 09/06/2029
        (r"([A-Za-z]{3,9})\.?\s*[-/ ,]?\s*(\d{4})", "my"),                         # March 2030
        (r"(\d{4})-(\d{1,2})\b", "ym"),                                             # 2030-03
        (r"\b(\d{1,2})[/.](\d{4})", "my"),                                          # 03/2030
    ]
    for pat, order in patterns:
        m = re.search(pat, t)
        if not m:
            continue
        parts = dict(zip(order, m.groups()))
        month = parts["m"]
        month = int(month) if month.isdigit() else MONTHS.get(month[:3].upper())
        day = int(parts.get("d", 1))
        if month in MONTH_ABBR and 1 <= day <= 31:
            return f"{day:02d}{MONTH_ABBR[month]}{parts['y']}"
    return None


def clean_formula(s):
    if not s:
        return None
    s = re.sub(r"\^2\s*H", "D", s)  # deuterium written as superscript 2H
    s = s.replace("^", "")
    if re.search(r"[•·]", s):
        parts = re.split(r"\s*[•·]\s*", s)
    else:
        parts = s.split()  # e.g. "C10H8D5NO2 HCl"
    return " • ".join(re.sub(r"\s+", "", p) for p in parts if p.strip())


def parse_concentration(s):
    """'A 1 mg/ml (nominal, as free base) solution in methanol' -> '1 mg/ml in methanol (as free base)'."""
    if not s:
        return None
    m = re.search(r"([\d.]+)\s*(µg|μg|ug|mg|ng|g)\s*/\s*m?l", s, re.I)
    if not m:
        return None
    unit = m.group(2).lower().replace("μ", "µ").replace("ug", "µg")
    conc = f"{m.group(1)} {unit}/ml"
    solvent = search(r"\bin\s+([A-Za-z][A-Za-z0-9:,\- ]*?)\s*(?:\(|$)", s[m.end():])
    if solvent:
        conc += f" in {solvent.lower()}"
    if "free base" in s.lower():
        conc += " (as free base)"
    return conc


def parse_form(supplied):
    low = supplied.lower()
    for word, form in [("solution", "Solution"), ("solid", "Solid"), ("powder", "Solid"),
                       ("oil", "Oil"), ("liquid", "Liquid"), ("gas", "Gas")]:
        if word in low:
            return form
    return None


# --------------------------------------------------------------------------- COA parser

def parse_coa(pdf_path_or_file, filename):
    with pdfplumber.open(pdf_path_or_file) as pdf:
        lines = page_lines(pdf.pages[0])
        page_width = pdf.pages[0].width
    texts = [ln["text"] for ln in lines]
    text = "\n".join(texts)
    if not text.strip():
        raise ValueError("no text layer found (scanned PDF?)")

    rec = {}
    rec["Lot No."] = search(r"(?:Batch|Lot)\s*(?:No\.?|Number)\s*:?\s*([A-Za-z0-9-]+)", text)
    rec["Product No."] = search(
        r"(?:Item|Product|Catalog(?:ue)?|Cat\.)\s*(?:No\.?|Number|Code)\s*:?\s*([A-Za-z0-9-]+)", text)
    rec["CAS No."] = search(r"CAS\b[^\n]{0,30}?(\d{2,7}-\d{2}-\d)\b", text)

    formula = search(r"Molecular\s+formula\.?\s*:?\s*(.+?)(?:\s*\(salt\)|\s*\(free base\)|\s{3,}|$)", text)
    rec["Chemical Formula"] = clean_formula(formula)

    mw = (search(r"Molecular\s+weight\s*:?\s*([\d.]+)\s*\(salt\)", text)
          or search(r"(?:Formula|Molecular)\s+weight\s*:?\s*([\d.]+)", text))
    rec["Drug Mol. Wt."] = float(mw) if mw else None

    expiry = (search(r"Expir(?:y|ation)\s+date\s*:?\s*(.+?)(?:\s{3,}|\(|$)", text)
              or search(r"(?:Re-?test|Use[- ]by)\s+date\s*:?\s*(.+?)(?:\s{3,}|\(|$)", text))
    rec["Expiry Date"] = format_date(expiry)

    # Chemical name + synonyms
    product_name = search(r"Product\s+name\s*:?\s*(.+)$", text)
    synonyms = search(r"Synonyms?\s*:?\s*(.+)$", text)
    item_idx = next((i for i, t in enumerate(texts) if re.search(r"(Item|Product)\s*No", t, re.I)), None)
    heading_idx = next((i for i, t in enumerate(texts)
                        if re.search(r"(CERTIFICATE|CONFIRMATION)\s+of\s+ANALYSIS", t, re.I)), None)
    if product_name:  # e.g. LGC: "d,l-MDA-D5.HCl, 1mg/ml in Methanol (as free base)"
        chemical = re.split(r",\s*[\d.]+\s*\S*g/m?l", first_segment(product_name), flags=re.I)[0]
    elif heading_idx is not None and heading_idx + 1 < len(texts):  # Cayman ISO certificate
        chemical = re.sub(r"\s*ACCREDITED$", "", first_segment(texts[heading_idx + 1]))
    else:  # Cayman one-page CofA: name on line 1, formal name until "Item No."
        chemical = first_segment(texts[0])
        if not synonyms and item_idx and item_idx > 1:
            synonyms = ""
            for prev, ln in zip(lines[0:item_idx - 1], lines[1:item_idx]):
                wrapped = prev is not lines[0] and prev["x1"] > 0.85 * page_width
                synonyms += ("" if wrapped or not synonyms else "; ") + first_segment(ln["text"])
    rec["Chemical"] = plain(chemical)
    rec["Synonyms"] = plain(synonyms)

    # Manufacturer
    haystack = (text + " " + filename).lower()
    rec["Manufacturer"] = next((name for key, name in KNOWN_MANUFACTURERS if key in haystack), None) \
        or search(r"(?:Manufacturer|Producer|Manufactured by)\s*:?\s*(\S.+?)(?:\s{3,}|$)", text)

    # Form + concentration
    supplied = search(r"Supplied\s+as\s*:?\s*(.+)$", text)
    if supplied:
        rec["Form"] = parse_form(supplied)
        rec["Concentration"] = parse_concentration(supplied)
    elif product_name and parse_concentration(product_name):
        rec["Form"] = "Solution"
        rec["Concentration"] = parse_concentration(product_name)

    for col in ["Chemical", "Synonyms", "Lot No.", "Product No.", "Expiry Date", "CAS No.",
                "Manufacturer", "Form", "Concentration", "Chemical Formula", "Drug Mol. Wt."]:
        if rec.get(col) in (None, ""):
            rec[col] = NA
    return rec


# --------------------------------------------------------------------------- output

def safe_name(s):
    return re.sub(r"[^A-Za-z0-9.-]+", "-", str(s)).strip("-") or NA


def write_excel(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(COLUMNS)
    for rec in rows:
        ws.append([rec.get(col, "") for col in COLUMNS])

    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx, col in enumerate(COLUMNS, 1):
        longest = max([len(col)] + [len(str(r.get(col, ""))) for r in rows])
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = min(max(12, longest + 2), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Populate an inventory Excel sheet from a zip of COA PDFs.")
    ap.add_argument("zips", nargs="+", type=Path, help="zip file(s) containing COA PDFs")
    ap.add_argument("-o", "--outdir", type=Path,
                    help="output folder (default: ./<first zip name>_inventory)")
    args = ap.parse_args()

    outdir = args.outdir or Path.cwd() / f"{args.zips[0].stem.replace(' ', '_')}_inventory"
    coa_dir = outdir / "COA"
    coa_dir.mkdir(parents=True, exist_ok=True)

    rows, used_names, problems = [], set(), []
    for zpath in args.zips:
        with zipfile.ZipFile(zpath) as zf:
            members = [m for m in zf.infolist()
                       if not m.is_dir() and m.filename.lower().endswith(".pdf")
                       and "__MACOSX" not in m.filename and not Path(m.filename).name.startswith("._")]
            for m in members:
                orig = Path(m.filename).name
                with zf.open(m) as fh:
                    data = fh.read()
                try:
                    rec = parse_coa(io.BytesIO(data), orig)
                except Exception as e:  # keep going; the file is still copied
                    print(f"  ! {orig}: could not read ({e})", file=sys.stderr)
                    problems.append(orig)
                    rec = {col: NA for col in ["Chemical", "Synonyms", "Lot No.", "Product No.",
                                               "Expiry Date", "CAS No.", "Manufacturer", "Form",
                                               "Concentration", "Chemical Formula", "Drug Mol. Wt."]}

                stem = f"{safe_name(rec['Lot No.'])}_{safe_name(rec['Expiry Date'])}"
                new_name, n = f"{stem}.pdf", 2
                while new_name in used_names:
                    new_name, n = f"{stem}_{n}.pdf", n + 1
                used_names.add(new_name)
                (coa_dir / new_name).write_bytes(data)
                rows.append(rec)

                missing = [k for k in ("Chemical", "Lot No.", "Product No.", "Expiry Date") if rec[k] == NA]
                flag = f"   [NA: {', '.join(missing)}]" if missing else ""
                print(f"  {orig}  ->  {new_name}   {rec['Chemical']}{flag}")

    xlsx = outdir / "COA_inventory.xlsx"
    write_excel(rows, xlsx)
    print(f"\n{len(rows)} COA(s) processed")
    print(f"Excel : {xlsx}")
    print(f"PDFs  : {coa_dir}")
    if problems:
        print(f"Unreadable PDFs ({len(problems)}): {', '.join(problems)}", file=sys.stderr)


if __name__ == "__main__":
    main()
